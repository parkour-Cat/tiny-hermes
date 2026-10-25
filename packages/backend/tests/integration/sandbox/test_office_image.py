"""The sandbox can do office work (§1.27, v2.27): spreadsheets, documents,
slides, PDFs and charts, with Chinese, inside the same locked-down container
every sandbox is — no network, read-only root, the sandbox user."""

import asyncio
import uuid
from typing import Any

from tiny_hermes.sandbox.domain.container_policy import DEFAULT_PROFILE, container_config
from tiny_hermes.sandbox.infrastructure.docker_engine import DockerEngine

from .conftest import LABEL

OFFICE_ROUND_TRIP = r'''
import io, json, warnings
warnings.simplefilter("always")
from openpyxl import Workbook, load_workbook
from docx import Document
from pptx import Presentation
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.backends.backend_pdf import PdfPages
from pypdf import PdfReader
from pdfminer.high_level import extract_text

out = {}
wb = Workbook(); ws = wb.active
ws.append(["部门", "销售额"]); ws.append(["华东", 1200]); wb.save("t.xlsx")
out["xlsx"] = load_workbook("t.xlsx").active["A2"].value
out["pandas"] = int(pd.read_excel("t.xlsx")["销售额"].sum())
d = Document(); d.add_heading("会议纪要", 1); d.add_paragraph("待办：张三 周五前"); d.save("t.docx")
out["docx"] = [p.text for p in Document("t.docx").paragraphs]
p = Presentation(); s = p.slides.add_slide(p.slide_layouts[0])
s.shapes.title.text = "季度汇报"; p.save("t.pptx")
out["pptx"] = Presentation("t.pptx").slides[0].shapes.title.text
# No font settings here: an Agent will not know to make any, so the defaults
# must already draw Chinese.
fig, ax = plt.subplots()
ax.bar(["华东", "华北"], [1200, 800]); ax.set_title("各部门销售额")
with warnings.catch_warnings(record=True) as caught:
    with PdfPages("t.pdf") as pdf:
        pdf.savefig(fig)
out["missing_glyphs"] = [str(w.message) for w in caught if "missing from font" in str(w.message)]
fig.savefig("t.png")
out["pdf_pages"] = len(PdfReader("t.pdf").pages)
out["pdf_text"] = extract_text("t.pdf")
print(json.dumps(out, ensure_ascii=False))
'''


async def test_office_files_round_trip_in_a_policy_shaped_sandbox(
    docker_client: Any, image_digest: str
) -> None:
    engine = DockerEngine(docker_client, extra_labels={LABEL: "1"})
    run_id = uuid.uuid4()
    docker_client.volumes.create(name=f"tiny-hermes-data-{run_id}", labels={LABEL: "1"})
    config = container_config(
        digest=image_digest,
        profile=DEFAULT_PROFILE,
        run_id=run_id,
        instance_id=uuid.uuid4(),
        workspace_id=uuid.uuid4(),
        approved_digests=(image_digest,),
    )
    container = docker_client.containers.get(await engine.create(config))

    code, output = await asyncio.to_thread(
        container.exec_run, ["python3", "-c", OFFICE_ROUND_TRIP], workdir="/workspace/data"
    )

    assert code == 0, output.decode(errors="replace")[-2000:]
    import json  # noqa: PLC0415

    found = json.loads(output.decode().strip().splitlines()[-1])
    assert found["xlsx"] == "华东"
    assert found["pandas"] == 1200
    assert found["docx"] == ["会议纪要", "待办：张三 周五前"]
    assert found["pptx"] == "季度汇报"
    assert found["missing_glyphs"] == [], "Chinese in a default chart would come out as boxes"
    assert found["pdf_pages"] == 1
    assert "各部门销售额" in found["pdf_text"]


async def test_the_office_image_still_has_no_network_tools_or_package_manager(
    docker_client: Any, image_digest: str
) -> None:
    """More in the image is not more ways out of it."""
    found = await asyncio.to_thread(
        docker_client.containers.run,
        image_digest,
        ["bash", "-c", "for b in curl wget apt-get pip pip3 nc ssh; do command -v $b; done; true"],
        remove=True,
        labels={LABEL: "1"},
    )

    assert found.decode().strip() == ""
